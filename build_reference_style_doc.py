from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

SRC='算法训练中心功能说明.md'
OUT='算法训练中心功能说明书-开发交付版.docx'
doc=Document()
sec=doc.sections[0]
sec.top_margin=Inches(.72); sec.bottom_margin=Inches(.7)
sec.left_margin=Inches(.82); sec.right_margin=Inches(.82)

for name,size,bold,color in [('Normal',10,False,'222222'),('Title',24,True,'000000'),('Heading 1',16,True,'000000'),('Heading 2',13,True,'000000'),('Heading 3',11,True,'000000')]:
    st=doc.styles[name]
    st.font.name='Microsoft YaHei'
    st._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    st.font.size=Pt(size); st.font.bold=bold; st.font.color.rgb=RGBColor.from_string(color)
    st.paragraph_format.space_after=Pt(8)
    if name=='Normal':
        st.paragraph_format.line_spacing=1.35

def set_cell_shading(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd')
    shd.set(qn('w:fill'),fill); tcPr.append(shd)

lines=open(SRC,encoding='utf-8').read().splitlines()
for line in lines:
    x=line.strip()
    if not x: continue
    if x.startswith('# '):
        p=doc.add_paragraph(x[2:],style='Title'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        sub=doc.add_paragraph('功能、字段、权限、交互及统计口径说明')
        sub.alignment=WD_ALIGN_PARAGRAPH.CENTER
        sub.runs[0].font.size=Pt(11); sub.runs[0].font.color.rgb=RGBColor(90,90,90)
        continue
    if x.startswith('## '):
        doc.add_heading(x[3:],level=1); continue
    if x.startswith('### '):
        doc.add_heading(x[4:],level=2); continue
    p=doc.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    if x.startswith('【') and x.endswith('】'):
        r=p.add_run(x); r.bold=True; r.font.size=Pt(11)
    elif re.match(r'^\d+）',x):
        m=re.match(r'^(\d+）)(.*)$',x)
        r=p.add_run(m.group(1)); r.bold=True
        p.add_run(m.group(2))
    else:
        p.add_run(x)

header=sec.header.paragraphs[0]
header.text='算法训练中心功能说明书'
header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.name='Microsoft YaHei'; header.runs[0].font.size=Pt(8); header.runs[0].font.color.rgb=RGBColor(110,110,110)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
run=footer.add_run('内部开发交付文档')
run.font.name='Microsoft YaHei'; run.font.size=Pt(8); run.font.color.rgb=RGBColor(110,110,110)

doc.core_properties.title='算法训练中心功能说明书'
doc.core_properties.subject='功能、字段、权限、交互及统计口径说明'
doc.save(OUT)
print(OUT)
